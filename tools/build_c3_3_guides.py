"""Genera las guías C-3-3 alineadas con el laboratorio obligatorio de voz."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
HEADER = ROOT / "src" / "teoria_senales" / "assets" / "encabezado_institucional.png"
STUDENT_DIR = ROOT / "docs" / "guias"
TEACHER_DIR = ROOT / "docs" / "guias_docente"
BLUE = "004B87"
LIGHT_BLUE = "EAF3FB"
YELLOW = "FFCF00"
INK = "13263F"
MUTED = "506176"
AUTHORS = "Juan Esteban Rengifo Segura; Julián Felipe Montaña Aguilar"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def configure(doc: Document, running_title: str) -> None:
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.9)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.4)
    section.footer_distance = Cm(0.7)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, before, after in (
        ("Heading 1", 15, 14, 7),
        ("Heading 2", 12.5, 10, 5),
        ("Heading 3", 11, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLUE)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(HEADER), width=Cm(17.0))

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(running_title)
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)

    doc.core_properties.author = AUTHORS
    doc.core_properties.last_modified_by = AUTHORS
    doc.core_properties.comments = "Proyecto Teoría de Señales"


def title(doc: Document, main: str, subtitle: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(main)
    run.font.name = "Arial"
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(BLUE)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(14)
    run = paragraph.add_run(subtitle)
    run.font.name = "Arial"
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(MUTED)


def table(doc: Document, rows: list[tuple[str, ...]], widths: list[float] | None = None):
    result = doc.add_table(rows=1, cols=len(rows[0]))
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.autofit = False
    result.style = "Table Grid"
    for index, value in enumerate(rows[0]):
        cell = result.rows[0].cells[index]
        cell.text = value
        set_cell_shading(cell, BLUE)
        set_cell_margins(cell)
        for run in cell.paragraphs[0].runs:
            run.font.name = "Arial"
            run.font.size = Pt(9)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for row_values in rows[1:]:
        cells = result.add_row().cells
        for index, value in enumerate(row_values):
            cells[index].text = value
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[index])
            for paragraph in cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8.8)
    if widths:
        for row in result.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return result


def bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Cm(0.55)
    paragraph.paragraph_format.first_line_indent = Cm(-0.3)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run(text)


def numbered(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.left_indent = Cm(0.65)
    paragraph.paragraph_format.first_line_indent = Cm(-0.35)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.add_run(text)


def callout(doc: Document, label: str, text: str, fill: str = LIGHT_BLUE) -> None:
    item = doc.add_table(rows=1, cols=1)
    item.alignment = WD_TABLE_ALIGNMENT.CENTER
    item.autofit = False
    cell = item.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=130, start=160, bottom=130, end=160)
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(label + " ")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(BLUE)
    paragraph.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_common_intro(doc: Document, teacher: bool) -> None:
    doc.add_heading("1. Información general" if not teacher else "1. Identificación del ejercicio", level=1)
    rows = [
        ("Campo", "Información"),
        ("Código y título", "C-3-3. Convolución lineal de señales discretas y verificación mediante DFT/FFT"),
        ("Componente del PDA", "Componente 3. Transformada de Fourier de tiempo discreto"),
        ("Subcomponente", "3.2. Propiedades de la Transformada de Fourier de tiempo discreto"),
        ("Modalidad", "Individual, con cuatro pasos guiados, laboratorio obligatorio de voz y evaluación final"),
        ("Duración estimada", "70 minutos"),
    ]
    table(doc, rows, [4.3, 12.2])

    doc.add_heading("2. Resultado de Aprendizaje Esperado", level=1)
    doc.add_paragraph(
        "Determinar la transformada de Fourier de tiempo discreto de señales comunes utilizadas en ingeniería, "
        "relacionar los resultados teóricos con la simulación y aplicar la convolución a una señal real de voz."
    )


def build_student() -> Path:
    doc = Document()
    configure(doc, "Guía del estudiante C-3-3")
    title(doc, "Guía del estudiante C-3-3", "Convolución lineal, DFT/FFT y laboratorio piloto de voz")
    add_common_intro(doc, teacher=False)

    doc.add_heading("3. Antes de comenzar", level=1)
    for text in (
        "Usa preferiblemente Chrome o Edge actualizado y abre Moodle mediante una conexión segura HTTPS.",
        "Comprueba que tu computador tenga micrófono y permite su uso cuando el navegador lo solicite.",
        "Ten disponible una palabra o frase corta. La captura se limita a cinco segundos.",
        "Si el micrófono está bloqueado, podrás cargar un archivo corto de voz como contingencia.",
    ):
        bullet(doc, text)
    callout(
        doc,
        "Importante:",
        "el laboratorio de voz es obligatorio para desbloquear la evaluación, pero durante el piloto no suma ni resta puntos.",
    )

    doc.add_heading("4. Enunciado y fundamento", level=1)
    doc.add_paragraph("Trabajarás con x[n] = {1, 2, 1} y h[n] = {1, -1, 2}.")
    doc.add_heading("4.1 Formulación matemática que debe aplicar", level=2)
    doc.add_paragraph("y[n] = x[n] * h[n] = Σ x[m] h[n-m]")
    doc.add_paragraph("L_y = L_x + L_h - 1; para evitar solapamiento circular se requiere N ≥ L_y.")

    doc.add_heading("5. Ruta de la actividad", level=1)
    steps = [
        ("Paso 1", "Reconocer L_x, L_h y calcular L_y."),
        ("Paso 2", "Calcular la convolución lineal muestra por muestra."),
        ("Paso 3", "Verificar mediante DFT/FFT con N = 3 e identificar el solapamiento circular."),
        ("Paso 4", "Aplicar relleno con ceros y verificar mediante DFT/FFT con N = 5."),
        ("Paso 5", "Completar el laboratorio obligatorio de voz."),
        ("Paso 6", "Responder la evaluación automática y finalizar el intento."),
    ]
    table(doc, [("Momento", "Qué debes hacer"), *steps], [3.0, 13.5])

    doc.add_heading("6. Paso 5: laboratorio obligatorio de voz", level=1)
    for text in (
        "Lee el aviso de privacidad temporal y marca la autorización de procesamiento.",
        "Pulsa Grabar voz, acepta el permiso del micrófono y pronuncia una palabra o frase corta.",
        "Escucha la señal original y la señal convolucionada con h[n] = [1, -1, 2].",
        "Observa la comparación de las formas de onda y relaciona el cambio con la suma de convolución.",
        "Cuando aparezca Laboratorio completado, continúa a la evaluación.",
    ):
        numbered(doc, text)
    callout(
        doc,
        "Contingencia técnica:",
        "si Moodle o el navegador bloquea el micrófono, revisa el permiso. Si el problema continúa, carga un archivo corto de voz. "
        "La evaluación permanecerá bloqueada hasta procesar una de las dos fuentes. Si tampoco puedes cargar el archivo, informa al tutor.",
        fill="FFF6D6",
    )

    doc.add_heading("7. Privacidad y decisión del piloto", level=1)
    doc.add_paragraph(
        "En esta versión, el audio se procesa temporalmente en el navegador. No se guarda el archivo original ni el procesado en "
        "Moodle o en el estado SCORM. Solo se registran metadatos técnicos: vía utilizada, duración, frecuencia de muestreo, número "
        "de muestras, finalización y fecha/hora del procesamiento."
    )
    doc.add_paragraph("Con el tutor deben discutir estas preguntas:")
    for text in (
        "¿Existe una necesidad pedagógica real de volver a escuchar cada grabación?",
        "¿Quién tendría acceso a los archivos y durante cuánto tiempo?",
        "¿Qué consentimiento, mecanismo de eliminación y controles ofrece Moodle?",
        "¿La mejora en retroalimentación justifica el tratamiento adicional de la voz?",
    ):
        bullet(doc, text)

    doc.add_page_break()
    doc.add_heading("8. Evaluación", level=1)
    rows = [
        ("Evidencia", "Cantidad", "Puntos"),
        ("Verificaciones numéricas guiadas", "4", "12"),
        ("Selección múltiple aleatoria", "3", "6"),
        ("Verdadero/Falso aleatorio", "2", "2"),
        ("Laboratorio de voz", "1 obligatorio", "0"),
        ("Total calificable", "9 respuestas", "20"),
    ]
    table(doc, rows, [8.5, 4.0, 4.0])
    doc.add_paragraph("Nota = (puntos obtenidos / 20) × 5,0. La aprobación se registra desde 3,0.")

    doc.add_heading("9. Lista de verificación", level=1)
    for text in (
        "Completé los cuatro pasos matemáticos.",
        "Procesé una señal de voz y comparé el audio original con el convolucionado.",
        "Comprendí que el archivo de audio no queda almacenado en esta versión.",
        "Respondí las cinco preguntas conceptuales y finalicé el intento.",
    ):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.2)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.add_run("☐ " + text)

    path = STUDENT_DIR / "Guia_del_estudiante_C-3-3.docx"
    doc.save(path)
    return path


def build_teacher() -> Path:
    doc = Document()
    configure(doc, "Guía técnica del docente C-3-3")
    title(doc, "Guía técnica del docente C-3-3", "Convolución lineal, DFT/FFT y laboratorio piloto obligatorio de voz")
    add_common_intro(doc, teacher=True)

    doc.add_heading("3. Propósito pedagógico", level=1)
    doc.add_paragraph(
        "Relacionar la convolución lineal en el dominio del tiempo con la multiplicación espectral y trasladar el procedimiento desde "
        "secuencias sintéticas a una señal real de voz. La obligatoriedad permite obtener evidencia comparable de usabilidad y viabilidad "
        "del micrófono durante el piloto."
    )
    callout(
        doc,
        "Regla del piloto:",
        "el laboratorio es requisito de avance, pero tiene ponderación 0. La nota permanece en escala 0,0–5,0 y se aprueba desde 3,0.",
    )

    doc.add_heading("4. Solución matemática de referencia", level=1)
    rows = [
        ("Procedimiento", "Resultado", "Interpretación"),
        ("Longitud", "L_x = 3; L_h = 3; L_y = 5", "L_y = L_x + L_h - 1"),
        ("Convolución lineal", "[1, 1, 1, 3, 2]", "Cinco muestras"),
        ("DFT/FFT con N = 3", "[4, 3, 1]", "Convolución circular con solapamiento"),
        ("DFT/FFT con N = 5", "[1, 1, 1, 3, 2]", "Coincide con la convolución lineal"),
    ]
    table(doc, rows, [5.2, 5.0, 6.3])

    doc.add_heading("5. Secuencia metodológica y tiempo", level=1)
    rows = [
        ("Momento", "Tiempo", "Evidencia"),
        ("Pasos 1–4", "43 min", "Cuatro verificaciones numéricas"),
        ("Laboratorio de voz", "10 min", "Señal procesada y metadatos SCORM"),
        ("Evaluación y cierre", "17 min", "Cinco preguntas y nota 0–5"),
    ]
    table(doc, rows, [6.0, 3.0, 7.5])

    doc.add_heading("6. Contrato técnico del laboratorio de voz", level=1)
    rows = [
        ("Elemento", "Definición"),
        ("Fuente principal", "Micrófono mediante getUserMedia y MediaRecorder"),
        ("Contingencia", "Carga de un archivo corto de voz cuando el micrófono esté bloqueado"),
        ("Duración", "Máximo 5 segundos"),
        ("Procesamiento", "Conversión mono a 16 kHz y convolución con h[n] = [1, -1, 2]"),
        ("Salida", "Reproducción original/procesada y comparación de forma de onda"),
        ("Finalización", "La evaluación se desbloquea solo con state.voice.completed = true"),
        ("Calificación", "Interacción SCORM con weighting = 0; no altera los 20 puntos"),
    ]
    table(doc, rows, [5.0, 11.5])

    doc.add_heading("7. Datos y privacidad", level=1)
    doc.add_paragraph(
        "La versión piloto no persiste bytes de audio. El archivo original y el procesado existen únicamente durante la pantalla del "
        "laboratorio. El estado SCORM conserva: fuente utilizada, duración, frecuencia de muestreo, número de muestras, indicador de "
        "truncamiento, finalización y fecha/hora."
    )
    doc.add_paragraph("Antes de aprobar almacenamiento futuro, el equipo debe resolver:")
    for text in (
        "Finalidad pedagógica concreta y proporcionalidad de conservar la voz.",
        "Consentimiento específico y alternativa para quien no autorice almacenamiento.",
        "Responsables, perfiles de acceso, ubicación, cifrado y registro de consultas.",
        "Plazo de conservación, eliminación y procedimiento para solicitudes del estudiante.",
        "Capacidad real de Moodle y lineamientos institucionales aplicables.",
    ):
        bullet(doc, text)

    doc.add_page_break()
    doc.add_heading("8. Protocolo de prueba en Moodle", level=1)
    for text in (
        "Publicar el ZIP como actividad SCORM 1.2 en un curso de prueba con HTTPS.",
        "Ingresar con perfil de estudiante en Chrome y Edge actualizados.",
        "Confirmar que la evaluación esté bloqueada antes de procesar voz.",
        "Aceptar el micrófono, grabar, detener y comprobar ambos reproductores y la forma de onda.",
        "Cerrar y reabrir el intento; verificar que la finalización del laboratorio se conserve sin guardar audio.",
        "Repetir con permiso denegado y validar la contingencia mediante archivo de voz.",
        "Finalizar con notas 3,0 y menor de 3,0; comprobar passed/failed, tiempo e interacciones.",
        "Revisar desde el perfil docente que C3-3-VOZ tenga ponderación 0 y que la nota permanezca 0–5.",
    ):
        numbered(doc, text)

    doc.add_heading("9. Solución de problemas", level=1)
    rows = [
        ("Situación", "Acción docente"),
        ("No aparece permiso", "Confirmar HTTPS, permiso del sitio y micrófono seleccionado en el sistema."),
        ("Permiso denegado", "Restablecer el permiso; si persiste, usar la carga de archivo."),
        ("No se procesa el audio", "Usar WAV, MP3, M4A u otro formato reproducible por el navegador; repetir con archivo corto."),
        ("La evaluación sigue bloqueada", "Confirmar que aparezca el mensaje Laboratorio obligatorio completado."),
        ("No hay micrófono ni archivo", "Registrar la incidencia y definir una sesión asistida; no modificar manualmente la nota."),
    ]
    table(doc, rows, [5.0, 11.5])

    doc.add_heading("10. Evaluación y criterios", level=1)
    rows = [
        ("Criterio", "Evidencia", "Puntos"),
        ("C3_3_CR1", "Identifica L_x, L_h y L_y", "3"),
        ("C3_3_CR2", "Calcula la convolución lineal", "3"),
        ("C3_3_CR3", "Resultado DFT/FFT para N = 3", "3"),
        ("C3_3_CR4", "Resultado DFT/FFT para N = 5", "3"),
        ("C3_3_CR5", "Tres preguntas de selección múltiple", "6"),
        ("C3_3_CR6", "Dos preguntas de verdadero/falso", "2"),
        ("C3-3-VOZ", "Laboratorio obligatorio completado", "0"),
        ("Total", "9 respuestas calificables", "20"),
    ]
    table(doc, rows, [4.0, 9.5, 3.0])
    doc.add_paragraph("Nota = (puntos obtenidos / 20) × 5,0. Aprobación desde 3,0.")

    doc.add_heading("11. Banco conceptual", level=1)
    doc.add_paragraph(
        "El banco conserva 15 preguntas: nueve de selección múltiple (C3-3-P05 a C3-3-P13) y seis de verdadero/falso "
        "(C3-3-P14 a C3-3-P19). En cada intento se seleccionan tres de selección múltiple y dos de verdadero/falso. "
        "Las claves y criterios permanecen definidos en el paquete y no cambian por la introducción de voz."
    )

    doc.add_heading("12. Criterios de aceptación del piloto", level=1)
    for text in (
        "Al menos 95 % de los intentos de prueba completan la captura o la contingencia sin asistencia técnica.",
        "El 100 % de los intentos mantiene la nota en escala 0–5 y el laboratorio con ponderación 0.",
        "Ningún intento persiste archivos de audio en el estado SCORM.",
        "Se documentan navegador, dispositivo, bloqueo presentado y sugerencia de cada tutor.",
        "La decisión sobre almacenamiento queda registrada antes de cambiar la política del piloto.",
    ):
        bullet(doc, text)

    path = TEACHER_DIR / "Guia_tecnica_del_docente_C-3-3.docx"
    doc.save(path)
    return path


def main() -> None:
    STUDENT_DIR.mkdir(parents=True, exist_ok=True)
    TEACHER_DIR.mkdir(parents=True, exist_ok=True)
    for path in (build_student(), build_teacher()):
        print(path.relative_to(ROOT))


if __name__ == "__main__":
    main()

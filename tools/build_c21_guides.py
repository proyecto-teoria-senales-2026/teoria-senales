from pathlib import Path
from copy import deepcopy

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "word_actualizadas"
STUDENT_TEMPLATE = OUT / "Guia_del_estudiante_C-3-3(2).docx"
TEACHER_TEMPLATE = OUT / "Guia_tecnica_del_docente_C-3-3(2).docx"
BLUE = "004B87"
LIGHT_BLUE = "EAF3FB"
PALE = "F4F7FA"
GRID = "D9D9D9"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(table):
    tbl_pr = table._tbl.tblPr
    node = tbl_pr.find(qn("w:tblBorders"))
    if node is None:
        node = OxmlElement("w:tblBorders")
        tbl_pr.append(node)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "6")
        item.set(qn("w:color"), GRID)
        node.append(item)


def set_cell_margin(cell, top=120, start=140, bottom=120, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = text
        shade(cell, BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = str(value)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                shade(cells[index], PALE)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    for row in table.rows:
        for cell in row.cells:
            set_cell_margin(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.08
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(9.5)
    borders(table)
    # Keep a minimal separation without forcing isolated content onto a new page.
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(2)
    spacer.add_run().font.size = Pt(2)
    return table


def configure(doc, title, subtitle):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    styles["Normal"].paragraph_format.space_after = Pt(5)
    styles["Normal"].paragraph_format.line_spacing = 1.08
    for name, size in (("Title", 24), ("Heading 1", 15), ("Heading 2", 12)):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
    title_ppr = styles["Title"].element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    styles["Heading 1"].paragraph_format.space_before = Pt(14)
    styles["Heading 1"].paragraph_format.space_after = Pt(6)
    styles["Heading 1"].paragraph_format.keep_with_next = True
    styles["Heading 2"].paragraph_format.space_before = Pt(9)
    styles["Heading 2"].paragraph_format.space_after = Pt(4)
    p = doc.add_paragraph(style="Title")
    p.add_run(title)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)
    p.paragraph_format.space_after = Pt(4)
    p = doc.add_paragraph(subtitle)
    p.runs[0].font.size = Pt(12)
    p.runs[0].font.bold = True
    p.paragraph_format.space_after = Pt(14)


def institutional_document(template, footer_text):
    """Reuse the approved institutional header and footer artwork."""
    doc = Document(template)
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    section = doc.sections[0]
    # Define the same institutional artwork for odd and even pages. This avoids
    # LibreOffice dropping the header on alternating pages during PDF export.
    even_odd = doc.settings._element.find(qn("w:evenAndOddHeaders"))
    if even_odd is None:
        even_odd = OxmlElement("w:evenAndOddHeaders")
        doc.settings._element.append(even_odd)
    else:
        even_odd.attrib.clear()
    title_page = section._sectPr.find(qn("w:titlePg"))
    if title_page is not None:
        section._sectPr.remove(title_page)
    default_header = section._sectPr.find(f"{qn('w:headerReference')}[@{qn('w:type')}='default']")
    even_header = section._sectPr.find(f"{qn('w:headerReference')}[@{qn('w:type')}='even']")
    if default_header is not None and even_header is None:
        even_header = deepcopy(default_header)
        even_header.set(qn("w:type"), "even")
        section._sectPr.insert(0, even_header)
    for paragraph in section.footer.paragraphs:
        paragraph.clear()
    footer = section.footer.paragraphs[0]
    footer.text = footer_text
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        run.font.name = "Arial"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(89, 105, 125)
    return doc


def institutional_title(doc, title, subtitle):
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    styles["Normal"].paragraph_format.space_after = Pt(5)
    styles["Normal"].paragraph_format.line_spacing = 1.08
    for name, size in (("Heading 1", 15), ("Heading 2", 12)):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 75, 135)
        style.font.bold = True
    styles["Heading 1"].paragraph_format.space_before = Pt(14)
    styles["Heading 1"].paragraph_format.space_after = Pt(6)
    styles["Heading 1"].paragraph_format.keep_with_next = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.font.name = "Arial"
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 75, 135)
    p.paragraph_format.space_after = Pt(4)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(subtitle)
    run.font.name = "Arial"
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor(89, 105, 125)
    p.paragraph_format.space_after = Pt(14)


def bullet(doc, text, numbered=False):
    doc.add_paragraph(text, style="List Number" if numbered else "List Bullet")


def build_student():
    doc = institutional_document(STUDENT_TEMPLATE, "Guía del estudiante C-2-1")
    institutional_title(doc, "Guía del estudiante C-2-1", "Transformada de Fourier continua")
    doc.add_heading("1. Información general", level=1)
    doc.add_paragraph("Esta guía acompaña el ejercicio C-2-1 antes de comenzar y durante sus cinco pasos. Explica qué debes calcular, cómo se distribuye la calificación y qué debes verificar antes de finalizar el intento en Moodle.")
    add_table(doc, ["Campo", "Información"], [
        ["Código y título", "C-2-1 Transformada de Fourier continua"],
        ["Modalidad", "Individual con cuatro pasos guiados y evaluación final"],
        ["Duración estimada", "60 minutos"],
        ["Calificación", "20 puntos convertidos a escala de 0 a 100"],
        ["Aprobación", "Desde 60 sobre 100"],
    ], [1.7, 5.0])
    doc.add_heading("2. Resultado de aprendizaje", level=1)
    doc.add_paragraph("Relacionar un pulso rectangular de tiempo continuo con su transformada de Fourier y analizar la función sinc, el valor en frecuencia cero, los ceros del espectro, la magnitud y la fase.")
    doc.add_heading("3. Antes de comenzar", level=1)
    bullet(doc, "Abre el ejercicio desde Moodle y confirma que aparezca tu nombre.")
    bullet(doc, "Consulta esta guía antes de iniciar; podrás abrirla nuevamente desde cada paso.")
    bullet(doc, "Escribe tus respuestas sin unidades cuando el campo solicite solamente un valor numérico.")
    bullet(doc, "Utiliza punto (.) como separador decimal. Ejemplo: 1.5.")
    bullet(doc, "La orientación ayuda a corregir el procedimiento, pero no completa la respuesta por ti.")
    doc.add_heading("4. Fundamentos", level=1)
    add_table(doc, ["Concepto", "Expresión", "Interpretación"], [
        ["Pulso rectangular", "x(t) = 1 para −1 ≤ t ≤ 1", "Tiene duración T = 2 s"],
        ["Transformada", "X(f) = ∫x(t)e^(−j2πft)dt", "Representa la señal en frecuencia"],
        ["Función sinc", "sinc(u) = sin(πu)/(πu)", "Describe la forma del espectro"],
        ["Valor X(0)", "X(0) = ∫x(t)dt", "Corresponde al área de la señal"],
        ["Ceros", "f = k/T, con k ≠ 0", "Su separación es inversa a T"],
    ], [1.25, 2.45, 3.0])
    doc.add_heading("5. Ruta de la actividad", level=1)
    add_table(doc, ["Paso", "Qué debes hacer", "Evidencia"], [
        ["1", "Identificar soporte y duración del pulso", "Tres valores"],
        ["2", "Seleccionar la transformada rectangular sinc", "Una selección"],
        ["3", "Calcular X(0) y el primer cero positivo", "Dos valores"],
        ["4", "Interpretar separación de ceros y fase", "Dos selecciones"],
        ["5", "Responder la evaluación automática y finalizar", "Cinco respuestas"],
    ], [0.7, 4.6, 1.4])
    doc.add_heading("6. Formulación de referencia", level=1)
    doc.add_paragraph("Paso 1: el soporte es [−1, 1] s y la duración es T = 2 s.")
    doc.add_paragraph("Paso 2: al sustituir T = 2, la transformada es X(f) = 2 sinc(2f).")
    doc.add_paragraph("Paso 3: X(0) = 2 y el primer cero positivo se ubica en f = 0.5 Hz.")
    doc.add_paragraph("Paso 4: al aumentar T, los ceros quedan más cercanos; cuando |X(f)| = 0, la fase no está definida.")
    doc.add_heading("7. Evaluación", level=1)
    add_table(doc, ["Evidencia", "Cantidad", "Puntos"], [
        ["Pasos guiados", "4", "12 máximo"],
        ["Selección múltiple aleatoria", "3", "6"],
        ["Verdadero o falso aleatorio", "2", "2"],
        ["Total", "9 evidencias", "20"],
    ], [3.6, 1.4, 1.7])
    doc.add_paragraph("Cada paso guiado asigna 3 puntos si aciertas en el primer intento y 2 puntos si corriges después. La nota se calcula como (puntos obtenidos / 20) × 100.")
    doc.add_heading("8. Lista de verificación", level=1)
    for item in ("Consulté la guía antes de iniciar.", "Completé los cuatro pasos guiados.", "Respondí las cinco preguntas de la evaluación.", "Revisé la nota en escala de 0 a 100.", "Usé Finalizar y salir y regresé correctamente a Moodle."):
        bullet(doc, item)
    path = OUT / "Guia_del_estudiante_C-2-1.docx"
    doc.save(path)
    return path


def build_teacher():
    doc = institutional_document(TEACHER_TEMPLATE, "Guía técnica del docente C-2-1")
    institutional_title(doc, "Guía técnica del docente C-2-1", "Transformada de Fourier continua")
    doc.add_heading("1. Identificación del ejercicio", level=1)
    add_table(doc, ["Campo", "Información"], [
        ["Código", "C-2-1"],
        ["Título", "Transformada de Fourier continua"],
        ["Modalidad", "Individual con cuatro pasos guiados y evaluación final"],
        ["Duración estimada", "60 minutos"],
        ["Estándar", "SCORM 1.2 para Moodle"],
    ], [1.7, 5.0])
    doc.add_heading("2. Resultado de aprendizaje", level=1)
    doc.add_paragraph("Relacionar un pulso rectangular de tiempo continuo con su transformada de Fourier y analizar la función sinc, el valor en frecuencia cero, los ceros del espectro, la magnitud y la fase.")
    doc.add_heading("3. Propósito pedagógico", level=1)
    doc.add_paragraph("Consolidar la relación rectangular sinc mediante una secuencia progresiva de identificación temporal, formulación de la transformada, cálculo espectral, interpretación y evaluación automática.")
    doc.add_heading("4. Solución de referencia", level=1)
    add_table(doc, ["Paso", "Resultado esperado", "Criterio"], [
        ["1", "Soporte [−1, 1] s; duración T = 2 s", "C2_1_CR1"],
        ["2", "X(f) = 2 sinc(2f)", "C2_1_CR2"],
        ["3", "X(0) = 2; primer cero positivo = 0.5 Hz", "C2_1_CR3"],
        ["4", "Ceros más cercanos; fase no definida en |X(f)| = 0", "C2_1_CR4"],
    ], [0.7, 4.7, 1.3])
    heading = doc.add_heading("5. Secuencia metodológica", level=1)
    heading.paragraph_format.page_break_before = True
    add_table(doc, ["Momento", "Tiempo", "Evidencia"], [
        ["Inicio y guía", "5 min", "Lectura de instrucciones"],
        ["Pasos 1 a 4", "40 min", "Cuatro verificaciones guiadas"],
        ["Evaluación y cierre", "15 min", "Cinco preguntas y nota 0 a 100"],
    ], [2.2, 1.1, 3.4])
    doc.add_heading("6. Evaluación y criterios", level=1)
    add_table(doc, ["Criterio", "Evidencia", "Puntos"], [
        ["C2_1_CR1", "Identifica soporte y duración", "3 primer intento o 2 después"],
        ["C2_1_CR2", "Obtiene la transformada rectangular sinc", "3 primer intento o 2 después"],
        ["C2_1_CR3", "Calcula X(0) y los ceros del espectro", "3 primer intento o 2 después"],
        ["C2_1_CR4", "Interpreta magnitud y fase", "3 primer intento o 2 después"],
        ["Banco conceptual", "Tres preguntas de selección múltiple", "6"],
        ["Banco conceptual", "Dos preguntas de verdadero o falso", "2"],
        ["Total", "Nueve evidencias calificables", "20 máximo"],
    ], [1.55, 3.75, 1.7])
    doc.add_paragraph("Nota = (puntos obtenidos / 20) × 100. Moodle registra el estado Pasado desde 60/100.")
    doc.add_heading("7. Comportamiento SCORM esperado", level=1)
    for item in ("Reconoce al estudiante desde Moodle.", "Mantiene la guía accesible antes del inicio y durante los cinco pasos.", "Presenta campos vacíos y opciones sin preselección.", "Registra intentos, interacciones, objetivos, progreso y tiempo.", "Conserva el avance al cerrar y reabrir el mismo intento.", "Ejecuta guardado, commit y LMSFinish antes de regresar a Moodle."):
        bullet(doc, item)
    doc.add_paragraph("Los campos numéricos utilizan punto (.) como separador decimal; por ejemplo, 1.5.")
    doc.add_heading("8. Protocolo de QA en Moodle", level=1)
    tests = (
        "Publicar el ZIP como SCORM 1.2 e iniciar un intento nuevo con perfil de estudiante.",
        "Confirmar nombre, guía inicial, enlace permanente y campos vacíos.",
        "Responder un paso incorrectamente, usar la orientación y corregirlo; revisar el registro de intentos.",
        "Cerrar y reabrir el intento para validar la persistencia.",
        "Completar una prueba de 20/20 y confirmar 100/100.",
        "Completar una combinación de 12/20 y confirmar 60/100 con estado Pasado.",
        "Completar una prueba menor de 12/20 y confirmar estado Fallido.",
        "Usar Finalizar y salir; comprobar el regreso a Moodle y la conservación de nota, estado y tiempo.",
    )
    for item in tests:
        bullet(doc, item, numbered=True)
    doc.add_heading("9. Banco conceptual", level=1)
    doc.add_paragraph("El banco contiene 15 preguntas: nueve de selección múltiple y seis de verdadero o falso. En cada intento se seleccionan tres de selección múltiple y dos de verdadero o falso mediante una semilla asociada al estudiante.")
    doc.add_heading("10. Criterios de aceptación", level=1)
    for item in ("El 100 % de los campos inicia vacío y sin respuestas reveladas.", "La guía se abre antes del inicio y desde cada paso.", "La nota se reporta en escala de 0 a 100 y se aprueba desde 60.", "Moodle conserva avance, interacciones, nota, estado y tiempo.", "Finalizar y salir devuelve al usuario fuera del reproductor sin perder los datos."):
        bullet(doc, item)
    path = OUT / "Guia_tecnica_del_docente_C-2-1.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    for created in (build_student(), build_teacher()):
        print(created)

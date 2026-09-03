from pathlib import Path
from docx import Document


ROOT = Path(__file__).resolve().parents[1] / "docs" / "word_actualizadas"


REPLACEMENTS = {
    "Observa la comparación de las formas de onda y relaciona el cambio con la suma de convolución.":
        "Observa por separado las tres gráficas: señal original, señal convolucionada y unión de ambas señales. Usa el control de instante para comparar sus amplitudes.",
    "Nota = (puntos obtenidos / 20) × 5,0. La aprobación se registra desde 3,0.":
        "Nota = (puntos obtenidos / 20) × 100. Durante el piloto, la aprobación se registra provisionalmente desde 60/100 y debe confirmarse en Moodle.",
    "Finalizar con notas 3,0 y menor de 3,0; comprobar passed/failed, tiempo e interacciones.":
        "Finalizar con notas 60 y menor de 60; comprobar passed/failed, tiempo e interacciones.",
    "Revisar desde el perfil docente que C3-3-VOZ tenga ponderación 0 y que la nota permanezca 0–5.":
        "Revisar desde el perfil docente que C3-3-VOZ tenga ponderación 0 y que la nota permanezca en escala 0–100.",
    "Nota = (puntos obtenidos / 20) × 5,0. Aprobación desde 3,0.":
        "Nota = (puntos obtenidos / 20) × 100. Aprobación provisional desde 60/100, sujeta a validación en Moodle.",
    "El 100 % de los intentos mantiene la nota en escala 0–5 y el laboratorio con ponderación 0.":
        "El 100 % de los intentos mantiene la nota en escala 0–100 y el laboratorio con ponderación 0.",
}


CELL_REPLACEMENTS = {
    "Regla del piloto: el laboratorio es requisito de avance, pero tiene ponderación 0. La nota permanece en escala 0,0–5,0 y se aprueba desde 3,0.":
        "Regla del piloto: el laboratorio es requisito de avance, pero tiene ponderación 0. La nota se reporta en escala 0–100 y se aprueba provisionalmente desde 60.",
    "Evaluación y cierre / 17 min / Cinco preguntas y nota 0–5":
        "Evaluación y cierre / 17 min / Cinco preguntas y nota 0–100",
    "Reproducción original/procesada y comparación de forma de onda":
        "Tres gráficas separadas: original, convolucionada y unión de ambas señales",
}


def iter_paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def replace_text(paragraph, old, new):
    if old not in paragraph.text:
        return
    for run in paragraph.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            return
    paragraph.text = paragraph.text.replace(old, new)


def update(path: Path):
    document = Document(path)
    for paragraph in iter_paragraphs(document):
        for old, new in REPLACEMENTS.items():
            replace_text(paragraph, old, new)
    for table in document.tables:
        for row in table.rows:
            joined = " / ".join(cell.text for cell in row.cells)
            for old, new in CELL_REPLACEMENTS.items():
                if old == joined:
                    values = new.split(" / ")
                    for cell, value in zip(row.cells, values):
                        cell.text = value
            for cell in row.cells:
                replace_text(cell.paragraphs[0], "Reproducción original/procesada y comparación de forma de onda", CELL_REPLACEMENTS["Reproducción original/procesada y comparación de forma de onda"])

    # Add the new controls to the procedure without altering the institutional layout.
    marker = "Escucha la señal original y la señal convolucionada con h[n] = [1, -1, 2]."
    for paragraph in document.paragraphs:
        if paragraph.text == marker:
            paragraph.text = marker + " Ajusta ganancia, separación/retardo y decaimiento para explorar su efecto."
        if paragraph.text == "Aceptar el micrófono, grabar, detener y comprobar ambos reproductores y la forma de onda.":
            paragraph.text = "Aceptar el micrófono, grabar y detener. Comprobar los reproductores, las tres gráficas y los controles de ganancia, separación/retardo, decaimiento e instante."

    # Remove orphaned explicit page breaks that would create a blank page.
    for paragraph in list(document.paragraphs):
        if not paragraph.text.strip() and "w:type=\"page\"" in paragraph._p.xml:
            paragraph._p.getparent().remove(paragraph._p)

    document.save(path)


for guide in ROOT.glob("*.docx"):
    update(guide)
